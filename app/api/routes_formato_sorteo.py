"""
Router FastAPI para generar el formato oficial de sorteo (Excel) a partir
de los CSV exportados por MZL (Resultados y Pruebas).

Cómo integrarlo en el proyecto:
1. Copia este archivo y `mapeo_formato_sorteo.py` a `app/routers/`.
2. Coloca la plantilla en blanco (el .xls convertido a .xlsx) en
   `app/static/plantillas/formato_sorteo_blanco.xlsx`. Debe conservar
   exactamente la misma estructura de celdas que se documentó en
   `mapeo_formato_sorteo.py` (si cambia el diseño del formato, ese es el
   único archivo que hay que ajustar).
3. En `main.py`:
       from app.routers import routes_formato_sorteo
       app.include_router(routes_formato_sorteo.router)
"""
import csv
import io
from pathlib import Path

from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook

from .mapeo_formato_sorteo import procesar_resultados, procesar_pruebas

router = APIRouter(prefix="/formato-sorteo", tags=["formato-sorteo"])

PLANTILLA_PATH = Path(__file__).resolve().parent.parent / "static" / "plantillas" / "formato_sorteo_blanco.xlsx"
NOMBRE_HOJA = "Hoja1"


def _leer_csv(contenido: bytes) -> list[dict]:
    texto = contenido.decode("utf-8-sig")  # utf-8-sig quita el BOM que exporta MZL
    return list(csv.DictReader(io.StringIO(texto)))


@router.post("/generar")
async def generar_formato_sorteo(
    resultados_csv: UploadFile = File(..., description="CSV de Resultados del sorteo"),
    pruebas_csv: UploadFile | None = File(None, description="CSV de Pruebas (opcional)"),
    numero_sorteo: str | None = Form(None, description="Sobrescribe la celda L19 (número de sorteo)"),
    fecha_texto: str | None = Form(None, description="Sobrescribe la celda P19 (fecha en texto)"),
):
    if not PLANTILLA_PATH.exists():
        raise HTTPException(
            status_code=500,
            detail=f"No se encontró la plantilla en {PLANTILLA_PATH}. Cópiala antes de usar este endpoint.",
        )

    wb = load_workbook(PLANTILLA_PATH)
    ws = wb[NOMBRE_HOJA]

    filas_resultados = _leer_csv(await resultados_csv.read())
    faltantes_resultados = procesar_resultados(ws, filas_resultados)

    faltantes_pruebas = []
    if pruebas_csv is not None:
        filas_pruebas = _leer_csv(await pruebas_csv.read())
        faltantes_pruebas = procesar_pruebas(ws, filas_pruebas)

    if numero_sorteo:
        ws["L19"] = int(numero_sorteo) if numero_sorteo.isdigit() else numero_sorteo
    if fecha_texto:
        ws["P19"] = fecha_texto

    if faltantes_resultados or faltantes_pruebas:
        # No bloqueamos la generación: devolvemos el archivo igual, pero avisamos
        # por encabezado cuáles premios no calzaron con el formato, para que el
        # usuario los revise manualmente en vez de perderlos en silencio.
        aviso = "; ".join(faltantes_resultados + faltantes_pruebas)
    else:
        aviso = ""

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)

    nombre_archivo = f"Formato_Sorteo_{numero_sorteo or 'generado'}.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{nombre_archivo}"'}
    if aviso:
        headers["X-Premios-No-Ubicados"] = aviso

    return StreamingResponse(
        salida,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )